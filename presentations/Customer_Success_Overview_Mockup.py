"""
Script to generate Customer Success Overview PowerPoint Presentation
This script creates a PowerPoint with the mockup image included.

To run this script:
1. Install python-pptx: pip install python-pptx
2. Place the mockup image in the same directory as this script
3. Run: python Customer_Success_Overview_Mockup.py
"""

from pptx import Presentation
from pptx.util import Inches

def create_powerpoint():
    """Creates a PowerPoint presentation with the Customer Success Overview mockup"""
    
    # Create a presentation object
    prs = Presentation()
    
    # Set slide width and height (16:9 aspect ratio)
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(5.625)
    
    # Add a blank slide
    blank_slide_layout = prs.slide_layouts[6]  # 6 is typically the blank layout
    slide = prs.slides.add_slide(blank_slide_layout)
    
    # Add the mockup image to the slide
    # The image should fill the entire slide
    img_path = "V4_3.3.1.1_Customer_Success_Overview_Importe_menu.png"
    
    try:
        # Add image to cover the entire slide
        left = Inches(0)
        top = Inches(0)
        width = prs.slide_width
        height = prs.slide_height
        
        pic = slide.shapes.add_picture(img_path, left, top, width=width, height=height)
        print(f"Image added successfully: {img_path}")
    except FileNotFoundError:
        print(f"Error: Image file not found: {img_path}")
        print("Please ensure the image file is in the same directory as this script.")
        return False
    except Exception as e:
        print(f"Error adding image: {e}")
        return False
    
    # Save the presentation
    output_filename = "Customer_Success_Overview_Mockup.pptx"
    prs.save(output_filename)
    print(f"PowerPoint created successfully: {output_filename}")
    return True

if __name__ == "__main__":
    create_powerpoint()
