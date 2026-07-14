#!/usr/bin/env python3
"""
Script to create a PowerPoint presentation with mockup images
Author: Customer Success Team
Date: 2026-07-14
Description: Creates a PowerPoint presentation with three mockup images in specified order
"""

try:
    from pptx import Presentation
    from pptx.util import Inches
    print("python-pptx module loaded successfully")
except ImportError as e:
    print(f"Error importing python-pptx: {e}")
    print("Please install python-pptx: pip install python-pptx")
    exit(1)

def create_presentation():
    """Create PowerPoint presentation with three mockup images"""
    
    # Create a presentation object
    print("Creating presentation object...")
    prs = Presentation()
    
    # Set slide dimensions (16:9 aspect ratio)
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(5.625)
    print(f"Slide dimensions set: {prs.slide_width} x {prs.slide_height}")
    
    # List of images in the specified order
    # Slide 1: V4_3.3.1.1_Customer_Success_Overview_Importe_menu.png
    # Slide 2: Prueba_2.png
    # Slide 3: Prueba_3.png
    images = [
        "V4_3.3.1.1_Customer_Success_Overview_Importe_menu.png",
        "Prueba_2.png",
        "Prueba_3.png"
    ]
    
    print(f"\nAdding {len(images)} slides...")
    
    # Add each image as a slide
    for idx, img_name in enumerate(images, 1):
        print(f"\nProcessing slide {idx}: {img_name}")
        
        # Add a blank slide
        blank_slide_layout = prs.slide_layouts[6]  # 6 is typically blank layout
        slide = prs.slides.add_slide(blank_slide_layout)
        
        # Add the image to cover the entire slide
        left = Inches(0)
        top = Inches(0)
        width = prs.slide_width
        height = prs.slide_height
        
        try:
            pic = slide.shapes.add_picture(img_name, left, top, width=width, height=height)
            print(f"  ✓ Successfully added image to slide {idx}")
        except FileNotFoundError:
            print(f"  ✗ Error: Image file not found: {img_name}")
            # Add a text placeholder
            textbox = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(1))
            text_frame = textbox.text_frame
            text_frame.text = f"Image not found: {img_name}"
        except Exception as e:
            print(f"  ✗ Error adding image {img_name}: {str(e)}")
    
    # Save the presentation
    output_file = "Customer_Success_Mockups_Presentation.pptx"
    print(f"\nSaving presentation as: {output_file}")
    prs.save(output_file)
    
    print(f"\n{'='*60}")
    print(f"SUCCESS! Presentation created successfully!")
    print(f"File: {output_file}")
    print(f"Total slides: {len(prs.slides)}")
    print(f"{'='*60}")
    
    return output_file

if __name__ == "__main__":
    create_presentation()
